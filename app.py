import streamlit as st
import google.generativeai as genai
from docx import Document
from docx.shared import Pt
from io import BytesIO
import re

# --- CONFIGURACIÓN DE IA ---
MODELO = 'gemini-2.0-flash' # O gemini-1.5-flash

try:
    api_key = st.secrets["GEMINI_API_KEY"]
    genai.configure(api_key=api_key)
    model = genai.GenerativeModel(MODELO)
except Exception as e:
    st.error("Error: Configura GEMINI_API_KEY en st.secrets")

# --- FUNCIONES DE PROCESAMIENTO ---
def extraer_anexos_estrictos(archivo):
    doc = Document(archivo)
    texto_completo = []
    for p in doc.paragraphs:
        if p.text.strip(): texto_completo.append(p.text)
    for t in doc.tables:
        for r in t.rows:
            texto_completo.append(" | ".join(c.text.strip() for c in r.cells))
    return "\n".join(texto_completo)

def exportar_word(texto):
    doc = Document()
    for line in texto.split('\n'):
        p = doc.add_paragraph(line)
        p.style.font.name = 'Arial'
        p.style.font.size = Pt(10)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()

# --- INTERFAZ ---
st.set_page_config(page_title="Bot Licitador Pro", layout="wide")
st.title("⚖️ Asistente Licitador Pro: GTN")

if 'paso' not in st.session_state: st.session_state.paso = 1
if 'anexo_actual_idx' not in st.session_state: st.session_state.anexo_actual_idx = 0
if 'lista_anexos' not in st.session_state: st.session_state.lista_anexos = []
if 'datos_registrados' not in st.session_state: st.session_state.datos_registrados = {}

# PASO 1: CARGA Y ESCANEO INICIAL
if st.session_state.paso == 1:
    archivo = st.file_uploader("Cargar Bases Integradas (.docx)", type="docx")
    if archivo and st.button("Analizar Estructura de Anexos"):
        with st.spinner("Escaneando formatos en las bases..."):
            texto = extraer_anexos_estrictos(archivo)
            st.session_state.raw_text = texto
            
            # Identificación de todos los bloques que empiezan con ANEXO N°
            prompt_init = (
                "Analiza el texto y extrae una lista de todos los encabezados de anexos encontrados. "
                "Si un anexo tiene variantes (ej: Persona Jurídica / Persona Natural), lístalas como items separados. "
                "Responde únicamente con los nombres de los anexos separados por '|'."
            )
            res = model.generate_content([prompt_init, texto]).text
            st.session_state.lista_anexos = [a.strip() for a in res.split('|') if len(a.strip()) > 5]
            st.session_state.paso = 2
            st.rerun()

# PASO 2: IDENTIFICACIÓN DE VARIANTE (OBLIGATORIO)
elif st.session_state.paso == 2:
    idx = st.session_state.anexo_actual_idx
    if idx < len(st.session_state.lista_anexos):
        anexo_nombre = st.session_state.lista_anexos[idx]
        st.header(f"Identificación: {anexo_nombre}")
        
        col1, col2 = st.columns(2)
        with col1:
            st.info(f"¿Desea procesar este anexo?")
            if st.button("📝 Llenar este anexo"):
                st.session_state.variante_seleccionada = anexo_nombre
                st.session_state.paso = 3
                st.rerun()
        with col2:
            st.warning("Omitir este formato")
            if st.button("⏩ Saltar este anexo"):
                st.session_state.anexo_actual_idx += 1
                st.rerun()
    else:
        st.success("Se han recorrido todos los anexos detectados.")
        if st.button("Reiniciar"):
            st.session_state.clear()
            st.rerun()

# PASO 3: LLENADO DE ANEXOS PERSONALIZADO (PÁRRAFOS Y TABLAS)
elif st.session_state.paso == 3:
    anexo = st.session_state.variante_seleccionada
    st.header(f"Llenado de Anexo Personalizado: {anexo}")

    # 3.1 Extraer el párrafo y campos (Solo una vez por anexo)
    if 'campos_anexo' not in st.session_state:
        with st.spinner("Extrayendo texto legal..."):
            prompt_p = (
                f"Busca el contenido completo del '{anexo}'. "
                "Extrae los párrafos que contienen espacios para completar [......] o etiquetas [CONSIGNAR...]. "
                "Devuelve el párrafo completo y una lista de las etiquetas encontradas separadas por comas."
            )
            res_p = model.generate_content([prompt_p, st.session_state.raw_text]).text
            # Intentar separar texto de lista de campos
            st.session_state.texto_borrador = res_p
            encontrados = re.findall(r'\[[^\]]+\]', res_p)
            st.session_state.campos_anexo = [c for c in encontrados if any(char.isalpha() for char in c)]

    # Interfaz de llenado
    st.subheader("Borrador del Párrafo")
    st.markdown(f"*{st.session_state.texto_borrador[:500]}...*") # Previsualización corta
    
    with st.form("form_llenado"):
        respuestas = {}
        st.markdown("### Complete los espacios en blanco del párrafo:")
        for i, campo in enumerate(st.session_state.campos_anexo):
            # Lógica de datos repetidos: el bot te los pide pero puedes ver qué pusiste antes
            val_prev = st.session_state.datos_registrados.get(campo, "")
            respuestas[campo] = st.text_input(f"Dato para {campo}", value=val_prev, key=f"in_{anexo}_{i}")
        
        # 3.2 CONSULTA DE TABLA
        st.markdown("---")
        llenar_tabla = st.checkbox("He detectado una tabla en este anexo. ¿Deseas completarla?", value=False)
        datos_tabla = ""
        if llenar_tabla:
            datos_tabla = st.text_area("Ingrese los datos de la tabla (o pegue desde Excel):", placeholder="RUC: 20... \nDomicilio: ...")

        if st.form_submit_button("✅ Finalizar y Generar"):
            st.session_state.datos_registrados.update(respuestas)
            
            with st.spinner("Redactando versión final..."):
                final_prompt = (
                    f"Redacta el {anexo} completo. Reemplaza los campos {respuestas} en el texto. "
                    f"Si el usuario dio estos datos de tabla: '{datos_tabla}', inclúyelos en el formato de tabla original. "
                    "Entrega el texto limpio y formal listo para imprimir."
                )
                st.session_state.resultado_docx = model.generate_content(final_prompt).text
                st.session_state.paso = 4
                st.rerun()

# PASO 4: RESULTADO Y DESCARGA
elif st.session_state.paso == 4:
    st.success("Anexo generado correctamente.")
    st.text_area("Vista Previa Final:", st.session_state.resultado_docx, height=400)
    
    # Descarga
    doc_download = exportar_word(st.session_state.resultado_docx)
    st.download_button("⬇️ Descargar Anexo en Word", data=doc_download, file_name=f"{st.session_state.variante_seleccionada}.docx")
    
    st.markdown("---")
    if st.button("Siguiente Anexo ➡️"):
        st.session_state.anexo_actual_idx += 1
        # Limpieza para el próximo anexo
        for k in ['campos_anexo', 'texto_borrador', 'resultado_docx', 'variante_seleccionada']:
            if k in st.session_state: del st.session_state[k]
        st.session_state.paso = 2
        st.rerun()
